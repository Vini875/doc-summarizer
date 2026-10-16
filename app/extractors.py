"""Text extraction for the three accepted formats: .docx, .txt and .pdf."""
import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .config import ALLOWED_EXTENSIONS, MAX_FILE_BYTES


class ExtractionError(ValueError):
    """The file was rejected or its text could not be read."""


def validate_upload(filename: str, data: bytes) -> str:
    """Check extension, size and file signature. Returns the normalized extension."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError("Only Word (.docx), Text (.txt) and PDF (.pdf) files are accepted.")
    if len(data) == 0:
        raise ExtractionError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError("The file is larger than the 1 MB limit.")

    # Check the content matches the extension, so a renamed file can't slip through.
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ExtractionError("The file has a .pdf extension but is not a valid PDF.")
    if ext == ".docx" and not data.startswith(b"PK"):
        raise ExtractionError("The file has a .docx extension but is not a valid Word document.")
    if ext == ".txt" and b"\x00" in data:
        raise ExtractionError("The file has a .txt extension but contains binary data.")
    return ext


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("The text file encoding could not be read. Save it as UTF-8 and try again.")


def _extract_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for corrupt files
        raise ExtractionError("The Word document could not be opened.") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("The PDF is password protected.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise ExtractionError("The PDF could not be read.") from exc
    return "\n".join(p for p in pages if p.strip())


_EXTRACTORS = {".txt": _extract_txt, ".docx": _extract_docx, ".pdf": _extract_pdf}


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Validate the upload and return (extension, extracted text)."""
    ext = validate_upload(filename, data)
    text = _EXTRACTORS[ext](data).strip()
    if not text:
        hint = " Scanned (image-only) PDFs are not supported." if ext == ".pdf" else ""
        raise ExtractionError("No readable text was found in the document." + hint)
    return ext, text
