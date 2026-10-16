import io

import pytest
from docx import Document

from app.extractors import ExtractionError, extract_text


def test_txt_is_extracted():
    ext, text = extract_text("notes.txt", "Hello world.\nSecond line.".encode("utf-8"))
    assert ext == ".txt"
    assert "Second line." in text


def test_docx_paragraphs_and_tables_are_extracted():
    doc = Document()
    doc.add_paragraph("Quarterly revenue grew by ten percent.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "North"
    buf = io.BytesIO()
    doc.save(buf)

    ext, text = extract_text("report.DOCX", buf.getvalue())
    assert ext == ".docx"
    assert "Quarterly revenue grew" in text
    assert "Region | North" in text


@pytest.mark.parametrize("name", ["image.png", "old.doc", "sheet.xlsx", "noextension"])
def test_other_formats_are_rejected(name):
    with pytest.raises(ExtractionError, match="Only Word"):
        extract_text(name, b"some content")


def test_files_over_one_megabyte_are_rejected():
    with pytest.raises(ExtractionError, match="1 MB"):
        extract_text("big.txt", b"a" * (1024 * 1024 + 1))


def test_exactly_one_megabyte_is_allowed():
    ext, _ = extract_text("edge.txt", b"a " * (512 * 1024))
    assert ext == ".txt"


@pytest.mark.parametrize("name", ["fake.pdf", "fake.docx"])
def test_renamed_files_are_rejected(name):
    with pytest.raises(ExtractionError, match="not a valid"):
        extract_text(name, b"just plain text pretending")


def test_empty_file_is_rejected():
    with pytest.raises(ExtractionError, match="empty"):
        extract_text("empty.txt", b"")
